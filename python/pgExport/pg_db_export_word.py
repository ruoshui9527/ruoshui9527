import psycopg2
from docx import Document
import pandas as pd
from openpyxl import load_workbook

def sync_database_structure():
    conn_src = psycopg2.connect(
        dbname="devdb",
        user="pg",
        password="pwd",
        port="58378",
        host="127.0.0.1"
    )

    sel_table_info = conn_src.cursor()
    sel_table_info.execute("SELECT ns.nspname AS table_schema, col.table_name, col.column_name, col.data_type,"
                           " CASE WHEN col.character_maximum_length IS NOT NULL THEN col.character_maximum_length::text "
                           "WHEN col.numeric_precision IS NOT NULL THEN col.numeric_precision::text ELSE '' END AS column_length, "
                           "CASE WHEN col.is_nullable = 'NO' THEN '是' ELSE '否' END AS required_field, obj_description(cls.oid, 'pg_class') AS table_comment, "
                           "col_description(cls.oid, col.ordinal_position::int) AS column_comment "
                           "FROM information_schema.columns col JOIN pg_class cls ON cls.relname = col.table_name JOIN pg_namespace ns "
                           "ON ns.oid = cls.relnamespace WHERE ns.nspname = 'safecommon' AND col.table_name LIKE 'sf_%' "
                           "ORDER BY col.table_name, col.ordinal_position")

    doc = Document()
    all_table = doc.add_table(rows=0, cols=14)
    all_table.style = 'Table Grid'

    sel_table_infos = sel_table_info.fetchall()
    for sel_table_info in sel_table_infos:
        table_schema = sel_table_info[0]
        table_name = sel_table_info[1]
        table_desc = sel_table_info[6]
        col_name = sel_table_info[2]
        col_data_type = sel_table_info[3]
        col_length = sel_table_info[4]
        col_required_field = sel_table_info[5]
        col_description = sel_table_info[7]

        row_cells = all_table.add_row().cells
        row_cells[0].text = table_desc if table_desc else ""
        row_cells[1].text = table_name
        row_cells[2].text = col_description if col_description else ""
        row_cells[3].text = col_name

        if col_data_type == 'USER-DEFINED':
            col_data_type = 'datetime'
        elif col_length:
            col_data_type = f'{col_data_type}({col_length})'

        row_cells[4].text = col_data_type
        row_cells[5].text = '/'
        row_cells[6].text = col_required_field
        row_cells[7].text = ''
        row_cells[8].text = ''
        row_cells[9].text = '/'

    doc.save('table_structure.docx')

    conn_src.close()


if __name__ == '__main__':
    sync_database_structure()